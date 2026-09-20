"""Extract plain text from uploaded resumes (PDF, DOCX, TXT)."""
from __future__ import annotations

import io
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


def extract_text(data: bytes, filename: str) -> str:
    """Return the text content of a resume given its raw bytes and filename."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _read_pdf(data)
    elif ext == ".docx":
        text = _read_docx(data)
    elif ext in {".txt", ".md"}:
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX or TXT.")
    return _normalize(text)


def extract_text_from_path(path: str | Path) -> str:
    p = Path(path)
    return extract_text(p.read_bytes(), p.name)


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _normalize(text: str) -> str:
    """Trim trailing spaces and collapse runs of blank lines."""
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    out: list[str] = []
    for line in lines:
        if line == "" and out and out[-1] == "":
            continue
        out.append(line)
    return "\n".join(out).strip()
